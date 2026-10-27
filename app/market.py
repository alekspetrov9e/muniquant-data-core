"""Market data area: separate, source-traceable observations of aluminium market instruments.

A table file (one registered document version) is read by a versioned adapter into a replayable
build. Every instrument, price type and field stays a separate value in the source's own currency
and unit. Nothing is converted, combined or derived here: no spreads, all-in prices, curve scores
or signals. Those are downstream transformations.
"""
import csv
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from uuid import uuid5, NAMESPACE_URL
from zoneinfo import ZoneInfo
from jsonschema import Draft202012Validator, FormatChecker
from sqlalchemy import select
from .models import (Audit, Commodity, Document, Instrument, InstrumentObservation, MarketBuild, Run, Source, SourceAccess,
                     PRICE_FIELDS, PRICE_UNITS, QUANTITY_FIELDS)
from .pipeline import logical_hash, snapshot_bytes
from .services import record

PARSER_VERSION = 'market-adapters-1.0.0'
PIPELINE_VERSION = '1.0.0'
VERSION = '1.0.0'
SCHEMA = Path(__file__).parent / 'market-v1.schema.json'
MAX_OBSERVATIONS = 50_000
VALUE_FIELDS = (*PRICE_FIELDS, *QUANTITY_FIELDS)
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
TSV = 'text/tab-separated-values'
# Downstream transformations the scope keeps out of the Core, refused anywhere in a package,
# together with the ten fields no Core export may carry.
SIGNALS = {'cash_3m_signal', 'lme_shfe_signal', 'premium_spread_signal', 'curve_score', 'tightness_score',
           'expected_move', 'all_in_price'}

def decimal_text(amount):
    text = format(amount, 'f')
    text = text.rstrip('0').rstrip('.') if '.' in text else text
    return '0' if text in ('-0', '') else text

def exact(text, field, where):
    """An exact decimal as written: '.' for decimals, optional ',' thousands grouping, no exponent."""
    value = text.strip()
    if ',' in value:
        if not re.fullmatch(r'-?[0-9]{1,3}(?:,[0-9]{3})+(?:\.[0-9]+)?', value): raise ValueError(f'{where}: {field} has invalid thousands grouping')
        value = value.replace(',', '')
    if not re.fullmatch(r'-?[0-9]+(?:\.[0-9]+)?', value): raise ValueError(f'{where}: {field} {text!r} is not a number')
    amount = Decimal(value)
    # Size first: quantizing a value beyond the decimal context's precision would itself fail.
    if abs(amount) >= Decimal('1e24') or amount != amount.quantize(Decimal('0.000001')):
        raise ValueError(f'{where}: {field} exceeds 24 integer or 6 decimal digits')
    if field in QUANTITY_FIELDS and amount < 0: raise ValueError(f'{where}: {field} is negative')
    return decimal_text(amount)

def utc(text):
    return datetime.fromisoformat(text).astimezone(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ') if text else None

def cells(raw, media, spec):
    """The table as (locator label, number, text cells) rows, header rows included, for evidence references."""
    if spec['adapter'] == 'market-csv-v1':
        # Comma-separated, or tab-separated as table copies from vendor screens are.
        if media not in (TSV, 'text/csv'): raise ValueError('The market CSV adapter requires a text/csv or tab-separated document')
        label, delimiter = ('TSV row', '\t') if media == TSV else ('CSV row', ',')
        return [(label, number, [cell.strip() for cell in row]) for number, row in enumerate(csv.reader(io.StringIO(raw.decode('utf-8-sig')), delimiter=delimiter), 1)]
    if spec['adapter'] == 'market-xlsx-v1':
        if media != XLSX or not raw.startswith(b'PK'): raise ValueError('The market XLSX adapter requires an .xlsx document')
        from openpyxl import load_workbook
        try: book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, OSError) as exc: raise ValueError('Malformed XLSX document') from exc
        try:
            if spec.get('sheet') not in book.sheetnames: raise ValueError('Sheet not found; the workbook has: ' + ', '.join(book.sheetnames))
            def text(value):
                # Cached cell values: dates as ISO dates or timestamps, numbers in their shortest exact form.
                if value is None: return ''
                if isinstance(value, datetime): return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
                if isinstance(value, date): return value.isoformat()
                if isinstance(value, bool): return str(value)
                if isinstance(value, float): return decimal_text(Decimal(repr(value)))
                return str(value).strip()
            rows = [(f"Sheet {spec['sheet']} row", number, [text(value) for value in row]) for number, row in enumerate(book[spec['sheet']].iter_rows(values_only=True), 1)]
        finally: book.close()
        return rows
    if spec['adapter'] == 'market-lseg-json-v1':
        return lseg_cells(raw, media)
    raise ValueError('Unknown market adapter')

def lseg_cells(raw, media):
    """An LSEG historical-pricing download as delivered: a list of request parts, each with the API's
    own column headers and data rows. Numbers are read exactly as written in the JSON."""
    if media != 'application/json': raise ValueError('The LSEG JSON adapter requires an application/json document')
    try: parts = json.loads(raw.decode('utf-8'), parse_float=Decimal)
    except ValueError as exc: raise ValueError('Malformed JSON document') from exc
    if not isinstance(parts, list) or not parts: raise ValueError('Expected a list of LSEG response parts')
    def text(value):
        if value is None: return ''
        if isinstance(value, bool) or not isinstance(value, (str, int, Decimal)): raise ValueError('Unexpected value in an LSEG response: ' + type(value).__name__)
        return format(value, 'f') if isinstance(value, Decimal) else str(value).strip()
    header, rows = None, []
    for number, part in enumerate(parts, 1):
        body = part.get('raw') if isinstance(part, dict) else None
        if not isinstance(body, dict) or not isinstance(body.get('headers'), list) or not isinstance(body.get('data'), list):
            raise ValueError(f'Response part {number} has no headers and data')
        names = [item.get('name') if isinstance(item, dict) else None for item in body['headers']]
        if not all(isinstance(name, str) for name in names): raise ValueError(f'Response part {number} has an unnamed column')
        if header is None: header = names
        elif names != header: raise ValueError(f'Response part {number} has other columns than part 1')
        for index, row in enumerate(body['data'], 1):
            if not isinstance(row, list) or len(row) != len(header): raise ValueError(f'Response part {number} row {index} does not match the headers')
            rows.append((f'Response part {number} row', index, [text(value) for value in row]))
    return [('Response header', 1, header), *rows]

def locate(places):
    """Evidence locator for the (label, number) places one observation was read from."""
    labels = {}
    for label, number in places: labels.setdefault(label, []).append(number)
    return '; '.join(f"{label}{'s' if len(numbers) > 1 else ''} {', '.join(map(str, numbers))}" for label, numbers in labels.items())

def parse(raw, media, spec):
    """Read one table into logical observation rows; pure, so a frozen file always replays the same."""
    lines = cells(raw, media, spec)
    header_index = spec['header_row'] - 1
    if header_index >= len(lines): raise ValueError('The header row is beyond the end of the table')
    header = lines[header_index][2]
    def position(name):
        if name is None: return None
        # '#N' names the N-th column, for vendor copies whose column names repeat.
        if re.fullmatch(r'#[1-9][0-9]{0,2}', name):
            if int(name[1:]) > len(header): raise ValueError(f"Column {name} is beyond the {len(header)} columns of header row {spec['header_row']}")
            return int(name[1:]) - 1
        found = [index for index, cell in enumerate(header) if cell == name]
        if len(found) != 1: raise ValueError(f"Column {name!r} must appear exactly once in header row {spec['header_row']}")
        return found[0]
    date_at, time_at = position(spec['date_column']), position(spec.get('time_column'))
    instrument_at, contract_at = position(spec.get('instrument_column')), position(spec.get('contract_code_column'))
    month_at, prompt_at = position(spec.get('contract_month_column')), position(spec.get('prompt_date_column'))
    columns = [(position(column['column']), column) for column in spec['columns']]
    missing, zone = set(spec['missing_values']), ZoneInfo(spec['timezone'])
    published, retrieved = utc(spec.get('published_at')), utc(spec['retrieved_at'])
    groups = {}
    leave_out, left_out = spec.get('leave_out') or {}, set()
    statuses = spec.get('status_by_date') or {}
    for label, number, row in lines[header_index + 1:]:
        if not any(row): continue
        cell = lambda index: row[index] if index is not None and index < len(row) else ''
        where = f'{label} {number}'
        def day(text, what):
            try: return datetime.strptime(text, spec['date_format']).date()
            except ValueError as exc: raise ValueError(f"{where}: {what} {text!r} does not match {spec['date_format']}") from exc
        effective = day(cell(date_at), 'date')
        # Named rows the Core cannot hold are skipped, with the reason kept in the build's settings.
        if effective.isoformat() in leave_out:
            left_out.add(effective.isoformat()); continue
        moment = None
        if time_at is not None and cell(time_at) not in missing:
            try: local = datetime.strptime(cell(date_at) + ' ' + cell(time_at), spec['date_format'] + ' ' + spec['time_format'])
            except ValueError as exc: raise ValueError(f'{where}: time {cell(time_at)!r} does not match {spec["time_format"]}') from exc
            moment = local.replace(tzinfo=zone).astimezone(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
        code = None
        if instrument_at is not None:
            if cell(instrument_at) not in spec['instrument_map']: raise ValueError(f'{where}: {cell(instrument_at)!r} is not in instrument_map')
            code = spec['instrument_map'][cell(instrument_at)]
        contract = cell(contract_at) or None if contract_at is not None and cell(contract_at) not in missing else None
        month = None
        if month_at is not None and cell(month_at) not in missing:
            match = re.fullmatch(r'(\d{4})-?(0[1-9]|1[0-2])', cell(month_at))
            if not match: raise ValueError(f'{where}: contract month {cell(month_at)!r} is not YYYY-MM or YYYYMM')
            month = match[1] + '-' + match[2]
        prompt = day(cell(prompt_at), 'prompt date').isoformat() if prompt_at is not None and cell(prompt_at) not in missing else None
        for index, column in columns:
            text = cell(index)
            if text in missing: continue
            value = exact(text, column['field'], f"{where}, column {column['column']}")
            instrument = column['instrument'] or code
            key = (instrument, column['price_type'], effective.isoformat(), moment, contract)
            group = groups.setdefault(key, {'instrument': instrument, 'price_type': column['price_type'], 'effective_date': effective.isoformat(),
                                            'effective_at': moment, 'contract_code': contract, 'contract_month': month, 'prompt_date': prompt,
                                            'values': {}, 'source_series': {}, 'published_at': published, 'retrieved_at': retrieved,
                                            'data_status': statuses.get(effective.isoformat(), spec['data_status']), 'rows': [], 'columns': []})
            if column['field'] in group['values']: raise ValueError(f"{where}: a second {column['field']} for {instrument} {column['price_type']} on {effective}")
            if (group['contract_month'], group['prompt_date']) != (month, prompt): raise ValueError(f'{where}: contract month or prompt date differs for one observation')
            group['values'][column['field']] = value
            group['source_series'][column['field']] = column.get('source_series') or column['column']
            if (label, number) not in group['rows']: group['rows'].append((label, number))
            if column['column'] not in group['columns']: group['columns'].append(column['column'])
        if len(groups) > MAX_OBSERVATIONS: raise ValueError(f'At most {MAX_OBSERVATIONS} observations per build')
    unknown = sorted((set(leave_out) - left_out) | (set(statuses) - {group['effective_date'] for group in groups.values()}))
    if unknown: raise ValueError('Dates named in leave_out or status_by_date are not in the table: ' + ', '.join(unknown))
    if not groups: raise ValueError('No values found; check the header row, columns and missing-value markers')
    rows = []
    for group in groups.values():
        values = {field: Decimal(value) for field, value in group['values'].items()}
        if 'high' in values and 'low' in values and values['high'] < values['low']: raise ValueError(f"{locate(group['rows'][:1])}: high is below low")
        if 'bid' in values and 'ask' in values and values['bid'] > values['ask']: raise ValueError(f"{locate(group['rows'][:1])}: bid is above ask")
        locator = f"{locate(group['rows'])}; columns {', '.join(group['columns'])}"
        rows.append({**{k: v for k, v in group.items() if k not in ('rows', 'columns')}, 'evidence_reference': locator})
    return rows

def logical_row(obs, instrument):
    """A stored observation in the parser's row form, for replay and export checks."""
    return {'instrument': instrument.code, 'price_type': obs.price_type, 'effective_date': obs.effective_date.isoformat(),
            'effective_at': obs.effective_at, 'contract_code': obs.contract_code, 'contract_month': obs.contract_month,
            'prompt_date': obs.prompt_date.isoformat() if obs.prompt_date else None,
            'values': {field: decimal_text(getattr(obs, field)) for field in VALUE_FIELDS if getattr(obs, field) is not None},
            'source_series': obs.source_series, 'published_at': obs.published_at, 'retrieved_at': obs.retrieved_at,
            'data_status': obs.data_status, 'evidence_reference': obs.evidence_reference}

def observation_id(build_id, index): return str(uuid5(NAMESPACE_URL, f'muniquant-market:{build_id}:{index}'))

def register_instrument(s, data):
    if not s.get(Commodity, data.commodity): raise ValueError('Unknown commodity')
    if s.scalar(select(Instrument).where(Instrument.code == data.code)): raise ValueError('Instrument code already registered: ' + data.code)
    obj = Instrument(**data.model_dump()); s.add(obj); s.flush(); return obj

def check_row(instrument, row):
    if instrument.contract_code and row['contract_code'] not in (None, instrument.contract_code):
        raise ValueError(f"{row['evidence_reference']}: contract {row['contract_code']} is not {instrument.code}'s contract {instrument.contract_code}")
    if not instrument.quantity_unit and any(field in row['values'] for field in QUANTITY_FIELDS):
        raise ValueError(f'{instrument.code} has volume or open interest but no quantity_unit (lots, contracts or t)')

def import_table(s, storage, document_id, data, actor):
    """Read a registered table document into a build and its observations; a repeat is a verified duplicate."""
    doc = s.get(Document, document_id)
    if not doc: raise ValueError('Unknown document')
    spec = data.model_dump(mode='json')
    build_id = logical_hash({'document_id': doc.id, 'content_hash': doc.content_hash, 'parser_version': PARSER_VERSION,
                             'pipeline_version': PIPELINE_VERSION, 'specification': spec})
    existing = s.get(MarketBuild, build_id)
    if existing:
        replay(s, storage, build_id)
        return existing, True
    manifest = {'adapter': data.adapter, 'document_id': doc.id, 'build_id': build_id, 'content_hash': doc.content_hash, 'parser_version': PARSER_VERSION}
    try:
        rows = parse(snapshot_bytes(storage, doc), doc.media_type, spec)
        codes = {row['instrument'] for row in rows}
        instruments = {i.code: i for i in s.scalars(select(Instrument).where(Instrument.code.in_(codes)))}
        if codes - set(instruments): raise ValueError('Register these instruments first: ' + ', '.join(sorted(codes - set(instruments))))
        for row in rows: check_row(instruments[row['instrument']], row)
    except (ValueError, UnicodeError) as exc:
        # Failures stay observable; no partial observation is kept.
        s.rollback()
        s.add(Run(status='failed', pipeline_version=PIPELINE_VERSION, manifest={**manifest, 'actor': actor, 'error': str(exc)[:500]})); s.commit()
        raise ValueError(str(exc)) from exc
    build = MarketBuild(id=build_id, document_id=doc.id, content_hash=doc.content_hash, parser_version=PARSER_VERSION, pipeline_version=PIPELINE_VERSION,
                        specification=spec, logical_output=rows, output_hash=logical_hash(rows), actor=actor)
    s.add(build); s.flush()
    for index, row in enumerate(rows):
        s.add(InstrumentObservation(id=observation_id(build_id, index), instrument_id=instruments[row['instrument']].id, document_id=doc.id, build_id=build_id,
                                    price_type=row['price_type'], effective_date=date.fromisoformat(row['effective_date']), effective_at=row['effective_at'],
                                    contract_code=row['contract_code'], contract_month=row['contract_month'],
                                    prompt_date=date.fromisoformat(row['prompt_date']) if row['prompt_date'] else None,
                                    **{field: Decimal(value) for field, value in row['values'].items()}, published_at=row['published_at'],
                                    retrieved_at=row['retrieved_at'], data_status=row['data_status'], source_series=row['source_series'],
                                    evidence_reference=row['evidence_reference']))
    s.add(Run(status='success', pipeline_version=PIPELINE_VERSION, manifest={**manifest, 'output_hash': build.output_hash, 'observations': len(rows)}))
    s.add(Audit(actor=actor, action='market_import', record_id=build_id, detail={'document_id': doc.id, 'observations': len(rows)}))
    s.flush()
    return build, False

def replay(s, storage, build_id):
    """Re-read the frozen document and prove the build and its stored observations are unchanged."""
    build = s.get(MarketBuild, build_id)
    if not build: raise ValueError('Unknown market build')
    if build.parser_version != PARSER_VERSION or build.pipeline_version != PIPELINE_VERSION: raise ValueError('Unsupported frozen parser/pipeline version')
    doc = s.get(Document, build.document_id)
    if not doc or doc.content_hash != build.content_hash: raise ValueError('Build provenance changed')
    rows = parse(snapshot_bytes(storage, doc), doc.media_type, build.specification)
    if logical_hash(rows) != build.output_hash or rows != build.logical_output: raise ValueError('Frozen market build output is not reproducible')
    instruments = {i.id: i for i in s.scalars(select(Instrument))}
    stored = {obs.id: obs for obs in s.scalars(select(InstrumentObservation).where(InstrumentObservation.build_id == build_id))}
    if len(stored) != len(rows): raise ValueError('Stored market observations differ from the build')
    for index, row in enumerate(rows):
        obs = stored.get(observation_id(build_id, index))
        if not obs or obs.document_id != build.document_id or logical_row(obs, instruments[obs.instrument_id]) != row:
            raise ValueError('Stored market observation differs from the frozen document: ' + observation_id(build_id, index))
    return {'build_id': build.id, 'parser_version': build.parser_version, 'input_hash': build.content_hash, 'output_hash': build.output_hash,
            'equivalent': True, 'observation_count': len(rows)}

def latest_access(s, source_id):
    return s.scalar(select(SourceAccess).where(SourceAccess.source_id == source_id).order_by(SourceAccess.decided_at.desc(), SourceAccess.id.desc()))

def exportable(s, source):
    """Licensed data leaves the Core only on an explicit access decision that allows export."""
    if source.source_type == 'synthetic': return True
    policy = latest_access(s, source.id)
    return bool(policy and policy.status == 'permitted' and policy.export_allowed)

def weekdays_between(first, second):
    return sum(1 for offset in range(1, (second - first).days) if date.fromordinal(first.toordinal() + offset).weekday() < 5)

def inspect_market_quality(s, storage):
    """Machine-readable findings for the market data; series checks are counted, not listed per row."""
    from .quality import snapshot_problem
    findings = []
    def add(code, severity, kind, identifier, message, count=None, examples=None):
        finding = {'code': code, 'severity': severity, 'record_type': kind, 'record_id': identifier, 'message': message,
                   'remediation': {'BLOCK': 'Correct the source data or its registration before export', 'ERROR': 'Correct the invalid record',
                                   'WARN': 'Review the source file and document the decision', 'INFO': 'No action unless the source says otherwise'}[severity]}
        if count is not None: finding.update(count=count, examples=examples or [])
        findings.append(finding)
    instruments = {i.id: i for i in s.scalars(select(Instrument).order_by(Instrument.code))}
    builds = {b.id: b for b in s.scalars(select(MarketBuild).order_by(MarketBuild.id))}
    observations = s.scalars(select(InstrumentObservation).order_by(InstrumentObservation.id)).all()
    documents = {d.id: d for d in s.scalars(select(Document).where(Document.id.in_({b.document_id for b in builds.values()}))).all()}
    for instrument in instruments.values():
        if PRICE_UNITS.get(instrument.unit) != instrument.currency: add('INVALID_UNIT', 'BLOCK', 'instrument', instrument.code, 'Unit and currency are unknown or inconsistent')
    checked_sources = set()
    for doc in documents.values():
        problem = snapshot_problem(storage, doc)
        if problem: add(problem[0], 'BLOCK', 'document', doc.id, problem[1])
        source = s.get(Source, doc.source_id)
        if source.id in checked_sources: continue
        checked_sources.add(source.id)
        policy = latest_access(s, source.id)
        if (policy and policy.status != 'permitted') or source.access_status not in ('permitted', 'synthetic'):
            add('ACCESS_REVOKED', 'BLOCK', 'source', source.id, 'Latest access decision does not permit this source')
        elif not policy and source.source_type != 'synthetic':
            add('MISSING_ACCESS_AUDIT', 'WARN', 'source', source.id, 'Source has no audited access decision; its data is not exported')
        elif policy and not policy.export_allowed:
            add('EXPORT_WITHHELD', 'INFO', 'source', source.id, 'Access decision keeps this source inside the Core; exports leave its data out')
    for build in builds.values():
        dropped = build.specification.get('leave_out') or {}
        if dropped: add('ROWS_LEFT_OUT', 'INFO', 'build', build.id, 'Rows named in the import settings were not loaded: ' + '; '.join(f'{d}: {r}' for d, r in sorted(dropped.items()))[:500], len(dropped), sorted(dropped)[:5])
    unknown = Counter(obs.build_id for obs in observations if obs.published_at is None)
    for build_id, count in sorted(unknown.items()):
        add('PUBLICATION_TIME_UNKNOWN', 'INFO', 'build', build_id, 'The source does not state when these values were published', count)
    for obs in observations:
        build = builds.get(obs.build_id)
        if not build or obs.document_id != build.document_id: add('PROVENANCE_MISMATCH', 'BLOCK', 'observation', obs.id, 'Observation and build cite different documents')
        if not instruments[obs.instrument_id].quantity_unit and (obs.volume is not None or obs.open_interest is not None):
            add('MISSING_QUANTITY_UNIT', 'BLOCK', 'observation', obs.id, 'Volume or open interest without a quantity unit')
    for status, code in (('carried_forward', 'DECLARED_CARRIED_FORWARD'), ('provisional', 'DECLARED_PROVISIONAL')):
        declared = [obs.id for obs in observations if obs.data_status == status]
        if declared: add(code, 'INFO', 'observation', declared[0], f'The source marks {len(declared)} values as {status.replace("_", " ")}', len(declared), declared[:5])
    series = defaultdict(list)
    for obs in observations: series[(obs.instrument_id, obs.price_type, obs.contract_code)].append(obs)
    for (instrument_id, price_type, contract), items in sorted(series.items(), key=lambda item: (instruments[item[0][0]].code, item[0][1], item[0][2] or '')):
        instrument = instruments[instrument_id]
        name = '/'.join(filter(None, (instrument.code, price_type, contract)))
        by_moment = defaultdict(list)
        for obs in items: by_moment[(obs.effective_date, obs.effective_at or '')].append(obs)
        conflicts, duplicates = [], []
        for moment, same in sorted(by_moment.items()):
            if len(same) < 2: continue
            values = [logical_row(obs, instrument)['values'] for obs in same]
            shared = set.intersection(*(set(v) for v in values))
            (conflicts if any(len({v[f] for v in values}) > 1 for f in shared) else duplicates).append(moment[0].isoformat())
        if conflicts: add('CONFLICTING_OBSERVATIONS', 'WARN', 'series', name, 'Documents give different values for the same date; both are kept', len(conflicts), conflicts[:5])
        if duplicates: add('DUPLICATE_OBSERVATION', 'INFO', 'series', name, 'Several documents give the same values for a date', len(duplicates), duplicates[:5])
        weekend = sorted({obs.effective_date.isoformat() for obs in items if obs.effective_date.weekday() >= 5})
        if weekend: add('WEEKEND_DATE', 'WARN', 'series', name, 'Values dated on a Saturday or Sunday; check for padded days', len(weekend), weekend[:5])
        # Day-to-day checks compare neighbours within one file: padding and gaps belong to a delivery.
        pairs = []
        for document_id in sorted({obs.document_id for obs in items}):
            ordered = sorted((obs for obs in items if obs.document_id == document_id), key=lambda obs: (obs.effective_date, obs.effective_at or ''))
            pairs += zip(ordered, ordered[1:])
        repeats, gaps, jumps = [], [], []
        for previous, current in pairs:
            before, after = logical_row(previous, instrument)['values'], logical_row(current, instrument)['values']
            prices = [field for field in PRICE_FIELDS if field in before and field in after]
            # A padded day copies everything: changed volume or open interest means real trading.
            quantities_equal = all(before.get(f) == after.get(f) for f in QUANTITY_FIELDS)
            if prices and quantities_equal and all(before[f] == after[f] for f in prices) and previous.effective_date != current.effective_date:
                repeats.append(current.effective_date.isoformat())
            if weekdays_between(previous.effective_date, current.effective_date) > 7: gaps.append(current.effective_date.isoformat())
            if instrument.market != 'physical' and any(Decimal(before[f]) and abs(Decimal(after[f]) - Decimal(before[f])) > abs(Decimal(before[f])) / 5 for f in prices):
                jumps.append(current.effective_date.isoformat())
        if repeats: add('POSSIBLE_CARRY_FORWARD', 'WARN', 'series', name, 'Every value identical to the previous date; may be copied forward by the vendor', len(repeats), sorted(repeats)[:5])
        if gaps: add('SERIES_GAP', 'WARN', 'series', name, 'More than seven weekdays without a value', len(gaps), sorted(gaps)[:5])
        if jumps: add('SUSPICIOUS_VALUE_CHANGE', 'WARN', 'series', name, 'A price moved more than 20% from the previous date', len(jumps), sorted(jumps)[:5])
    severities = Counter(f['severity'] for f in findings)
    status = 'FAIL' if severities['BLOCK'] or severities['ERROR'] else 'WARN' if severities['WARN'] else 'PASS'
    return {'status': status, 'counts': dict(severities), 'findings': findings,
            'coverage': {'instruments': len(instruments), 'series': len(series), 'observations': len(observations), 'builds': len(builds), 'documents': len(documents)}}

def validate_market_package(package):
    from .contract import PROHIBITED
    def visit(value):
        if isinstance(value, dict):
            if (PROHIBITED | SIGNALS).intersection(value): raise ValueError('Analytical or downstream fields are prohibited')
            for child in value.values(): visit(child)
        elif isinstance(value, list):
            for child in value: visit(child)
    visit(package)
    Draft202012Validator(json.loads(SCHEMA.read_text()), format_checker=FormatChecker()).validate(package)
    def indexed(items, key):
        mapping = {item[key]: item for item in items}
        if len(mapping) != len(items): raise ValueError('Duplicate market package identity: ' + key)
        return mapping
    instruments = indexed(package['instruments'], 'code')
    sources = indexed(package['sources'], 'source_id')
    docs = indexed(package['documents'], 'document_id')
    builds = indexed(package['build_manifest'], 'build_id')
    indexed(package['observations'], 'observation_id')
    for doc in docs.values():
        if doc['source_id'] not in sources: raise ValueError('Document source is absent')
    for item in package['observations']:
        instrument, doc, build = instruments.get(item['instrument_code']), docs.get(item['document_id']), builds.get(item['build_id'])
        if not instrument or not doc or not build or build['document_id'] != doc['document_id'] or doc['content_hash'] != item['content_hash'] or doc['source_id'] != item['source_id']:
            raise ValueError('Market observation provenance join is invalid')
        if any(item[field] != instrument[field] for field in ('exchange', 'market', 'commodity', 'grade', 'region', 'tenor', 'currency', 'unit', 'quantity_unit')):
            raise ValueError('Market observation identity differs from its instrument')
        if all(item[field] is None for field in VALUE_FIELDS): raise ValueError('Market observation has no value')
    for build in builds.values():
        doc = docs.get(build['document_id'])
        if not doc or doc['content_hash'] != build['content_hash']: raise ValueError('Build manifest input document is absent or inconsistent')
    if logical_hash({k: v for k, v in package.items() if k != 'build_hash'}) != package['build_hash']: raise ValueError('Package logical hash mismatch')
    return package

def export_market_v1(s, storage, codes=None, start=None, end=None):
    """MarketDataPackage 1.0.0: the selected observations of exportable sources, each self-contained."""
    quality = inspect_market_quality(s, storage)
    if quality['status'] == 'FAIL': raise ValueError('Market data quality BLOCK/ERROR findings prevent export; inspect /api/market/quality')
    instruments = {i.id: i for i in s.scalars(select(Instrument))}
    if codes:
        unknown = set(codes) - {i.code for i in instruments.values()}
        if unknown: raise ValueError('Unknown instruments: ' + ', '.join(sorted(unknown)))
    query = select(InstrumentObservation)
    if start: query = query.where(InstrumentObservation.effective_date >= start)
    if end: query = query.where(InstrumentObservation.effective_date <= end)
    selected = [obs for obs in s.scalars(query) if not codes or instruments[obs.instrument_id].code in codes]
    documents = {doc_id: s.get(Document, doc_id) for doc_id in sorted({obs.document_id for obs in selected})}
    sources = {doc.source_id: s.get(Source, doc.source_id) for doc in documents.values()}
    allowed = {source_id for source_id, source in sources.items() if exportable(s, source)}
    withheld = Counter(documents[obs.document_id].source_id for obs in selected if documents[obs.document_id].source_id not in allowed)
    selected = [obs for obs in selected if documents[obs.document_id].source_id in allowed]
    selected.sort(key=lambda obs: (instruments[obs.instrument_id].code, obs.price_type, obs.effective_date, obs.effective_at or '', obs.contract_code or '', obs.id))
    codes_by_record = defaultdict(set)
    for finding in quality['findings']:
        if finding['severity'] == 'WARN': codes_by_record[finding['record_id']].add(finding['code'])
    builds = {}
    for build_id in sorted({obs.build_id for obs in selected}):
        replay(s, storage, build_id)
        builds[build_id] = s.get(MarketBuild, build_id)
    records = []
    for obs in selected:
        instrument, doc = instruments[obs.instrument_id], documents[obs.document_id]
        source = sources[doc.source_id]
        series = '/'.join(filter(None, (instrument.code, obs.price_type, obs.contract_code)))
        row = logical_row(obs, instrument)
        quality_codes = sorted(codes_by_record[series] | codes_by_record[obs.id] | codes_by_record[doc.id] | codes_by_record[source.id])
        records.append({'observation_id': obs.id, 'instrument_code': instrument.code, 'instrument_name': instrument.name, 'exchange': instrument.exchange,
                        'market': instrument.market, 'commodity': instrument.commodity, 'grade': instrument.grade, 'region': instrument.region,
                        'tenor': instrument.tenor, 'contract_code': obs.contract_code or instrument.contract_code,
                        'contract_month': obs.contract_month or instrument.contract_month, 'prompt_date': row['prompt_date'],
                        'price_type': obs.price_type, **{field: row['values'].get(field) for field in VALUE_FIELDS},
                        'currency': instrument.currency, 'unit': instrument.unit, 'quantity_unit': instrument.quantity_unit,
                        'effective_date': row['effective_date'], 'effective_at': obs.effective_at, 'published_at': obs.published_at,
                        'retrieved_at': obs.retrieved_at, 'data_status': obs.data_status, 'source_id': source.id, 'publisher': source.publisher,
                        'source_series': obs.source_series, 'document_id': doc.id, 'content_hash': doc.content_hash,
                        'evidence_reference': obs.evidence_reference, 'build_id': obs.build_id, 'quality_status': 'WARN' if quality_codes else 'PASS',
                        'quality_codes': quality_codes, 'synthetic': source.source_type == 'synthetic', 'parser_version': PARSER_VERSION,
                        'pipeline_version': PIPELINE_VERSION, 'package_version': VERSION})
    used = sorted({obs.instrument_id for obs in selected}, key=lambda identifier: instruments[identifier].code)
    package = {'package_version': VERSION, 'pipeline_version': PIPELINE_VERSION,
               'selection': {'instruments': sorted(codes) if codes else None, 'from': start.isoformat() if start else None, 'to': end.isoformat() if end else None},
               'instruments': [{key: value for key, value in record(instruments[i]).items() if key not in ('id', 'created_at')} for i in used],
               'observations': records,
               'sources': [{'source_id': source.id, 'name': source.name, 'publisher': source.publisher, 'source_type': source.source_type, 'url': source.url,
                            'access_notes': source.notes} for source_id, source in sorted(sources.items()) if source_id in allowed and any(documents[obs.document_id].source_id == source_id for obs in selected)],
               'documents': [{'document_id': doc.id, 'source_id': doc.source_id, 'title': doc.title, 'original_url': doc.original_url, 'media_type': doc.media_type,
                              'version': doc.version, 'published_at': doc.published_at, 'retrieved_at': doc.retrieved_at, 'content_hash': doc.content_hash}
                             for doc in documents.values() if any(obs.document_id == doc.id for obs in selected)],
               'build_manifest': [{'build_id': build.id, 'document_id': build.document_id, 'content_hash': build.content_hash, 'parser_version': build.parser_version,
                                   'pipeline_version': build.pipeline_version, 'output_hash': build.output_hash} for build in builds.values()],
               'withheld': [{'source_id': source_id, 'publisher': sources[source_id].publisher, 'observations': count} for source_id, count in sorted(withheld.items())]}
    package['build_hash'] = logical_hash(package)
    return validate_market_package(package)

def freeze_market_v1(s, storage, target, codes=None, start=None, end=None):
    """A frozen bundle: the package, its builds and instruments, and the exact document bytes."""
    package = export_market_v1(s, storage, codes, start, end)
    builds = [record(s.get(MarketBuild, item['build_id'])) for item in package['build_manifest']]
    recipe = {'package': package, 'builds': builds}
    recipe['recipe_hash'] = logical_hash(recipe)
    with zipfile.ZipFile(target, 'x', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('recipe.json', json.dumps(recipe, sort_keys=True, indent=2))
        for digest in sorted({doc['content_hash'] for doc in package['documents']}): archive.write(Path(storage) / digest, 'snapshots/' + digest)
    return package['build_hash']

def replay_market_v1(target):
    """Rebuild every exported observation from the frozen document bytes, without any database."""
    import hashlib
    with zipfile.ZipFile(target) as archive:
        if len({item.filename for item in archive.infolist()}) != len(archive.infolist()): raise ValueError('Duplicate archive entry')
        if sum(item.file_size for item in archive.infolist()) > 200_000_000: raise ValueError('Bundle exceeds 200 MB limit')
        recipe = json.loads(archive.read('recipe.json'))
        if logical_hash({k: v for k, v in recipe.items() if k != 'recipe_hash'}) != recipe['recipe_hash']: raise ValueError('Recipe integrity failure')
        package = validate_market_package(recipe['package'])
        docs = {doc['document_id']: doc for doc in package['documents']}
        manifests = {item['build_id']: item for item in package['build_manifest']}
        if set(manifests) != {build['id'] for build in recipe['builds']}: raise ValueError('Frozen build set differs')
        expected = {}
        for build in recipe['builds']:
            if build['parser_version'] != PARSER_VERSION or build['pipeline_version'] != PIPELINE_VERSION: raise ValueError('Unsupported frozen parser/pipeline version')
            if any(manifests[build['id']][field] != build[field] for field in ('document_id', 'content_hash', 'parser_version', 'pipeline_version', 'output_hash')): raise ValueError('Frozen build manifest differs')
            identity = logical_hash({field: build[field] for field in ('document_id', 'content_hash', 'parser_version', 'pipeline_version')} | {'specification': build['specification']})
            if identity != build['id']: raise ValueError('Frozen build identity differs')
            raw = archive.read('snapshots/' + build['content_hash'])
            if hashlib.sha256(raw).hexdigest() != build['content_hash']: raise ValueError('Frozen snapshot integrity failure')
            rows = parse(raw, docs[build['document_id']]['media_type'], build['specification'])
            if rows != build['logical_output'] or logical_hash(rows) != build['output_hash']: raise ValueError('Frozen market extraction differs')
            for index, row in enumerate(rows): expected[observation_id(build['id'], index)] = row
        for item in package['observations']:
            row = expected.get(item['observation_id'])
            if not row: raise ValueError('Exported observation has no frozen source row')
            if (item['instrument_code'], item['price_type'], item['effective_date'], item['effective_at'], item['prompt_date'], item['source_series'], item['evidence_reference'], item['data_status']) != \
               (row['instrument'], row['price_type'], row['effective_date'], row['effective_at'], row['prompt_date'], row['source_series'], row['evidence_reference'], row['data_status']) or \
               any(item[field] != row['values'].get(field) for field in VALUE_FIELDS):
                raise ValueError('Exported observation differs from its frozen source row')
        return {'equivalent': True, 'build_hash': package['build_hash'], 'builds': len(recipe['builds']), 'observations': len(package['observations'])}

def market_summary(s):
    """Per instrument: how many observations, which dates, and the latest values; for the screens."""
    from sqlalchemy import func
    spans = {row[0]: row[1:] for row in s.execute(select(InstrumentObservation.instrument_id, func.count(), func.min(InstrumentObservation.effective_date),
                                                         func.max(InstrumentObservation.effective_date)).group_by(InstrumentObservation.instrument_id))}
    result = []
    for instrument in s.scalars(select(Instrument).order_by(Instrument.code)):
        count, first, last = spans.get(instrument.id, (0, None, None))
        latest = [logical_row(obs, instrument) | {'document_id': obs.document_id} for obs in s.scalars(select(InstrumentObservation).where(
            InstrumentObservation.instrument_id == instrument.id, InstrumentObservation.effective_date == last).order_by(InstrumentObservation.price_type, InstrumentObservation.id))] if last else []
        result.append({**{k: v for k, v in record(instrument).items() if k != 'created_at'}, 'observations': count,
                       'first_date': first.isoformat() if first else None, 'last_date': last.isoformat() if last else None, 'latest': latest})
    return {'instruments': result}
